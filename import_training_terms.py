"""Import terminology mappings from training workbook into learned terminology."""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl

from engineering_terminology import EngineeringTerminology


def normalize(value: object) -> str:
    return str(value or "").strip()


def main() -> int:
    parser = argparse.ArgumentParser(description="Import acronym/term expansions from xlsx")
    parser.add_argument(
        "--xlsx",
        default="Training Materials/Caltrans_Bridge_Engineering_Master_Terminology_400.xlsx",
        help="Path to source workbook",
    )
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    if not xlsx_path.exists():
        print(f"ERROR: Workbook not found: {xlsx_path}")
        return 1

    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    ws = wb[wb.sheetnames[0]]

    headers = [normalize(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    try:
        term_idx = headers.index("Term")
        expanded_idx = headers.index("Expanded")
    except ValueError:
        print(f"ERROR: Expected headers not found. Found: {headers}")
        return 1

    existing = EngineeringTerminology._combined_terms()
    before_pairs = set()
    for primary, syns in existing.items():
        for s in syns:
            before_pairs.add((primary.lower(), s.lower()))

    total_rows = 0
    skipped_rows = 0
    added_rows = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        total_rows += 1
        term = normalize(row[term_idx] if term_idx < len(row) else "")
        expanded = normalize(row[expanded_idx] if expanded_idx < len(row) else "")

        if not term or not expanded:
            skipped_rows += 1
            continue
        if term.lower() == expanded.lower():
            skipped_rows += 1
            continue

        EngineeringTerminology.add_feedback_mapping(term, expanded)

    updated = EngineeringTerminology._combined_terms()
    after_pairs = set()
    for primary, syns in updated.items():
        for s in syns:
            after_pairs.add((primary.lower(), s.lower()))

    added_rows = len(after_pairs - before_pairs)

    print(f"Workbook: {xlsx_path}")
    print(f"Rows scanned: {total_rows}")
    print(f"Rows skipped: {skipped_rows}")
    print(f"New mappings added: {added_rows}")
    print("Sample expansions:")
    for key in ["LOTB", "CIDH", "ABUT", "AASHTO"]:
        print(f"  {key}: {EngineeringTerminology.get_all_synonyms(key)[:4]}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
